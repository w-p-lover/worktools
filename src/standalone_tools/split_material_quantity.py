#!/usr/bin/env python3
"""Split model/count pairs in an Excel material-detail column."""

from __future__ import annotations

import argparse
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


HEADER_NAME = "物料明细"
RESULT_SHEET_NAME = "拆分结果"
MODEL_PATTERN = r"[A-Za-z]{1,6}\d*(?:[-.][A-Za-z0-9]+)+"
PAIR_PATTERN = re.compile(
    rf"(?P<model>{MODEL_PATTERN})\s*(?:样机|定频)?\s*"
    r"(?P<count>\d+|[零〇一二两三四五六七八九十百]+)\s*台",
    re.IGNORECASE,
)
MODEL_ONLY_PATTERN = re.compile(MODEL_PATTERN, re.IGNORECASE)

CHINESE_DIGITS = {
    "零": 0,
    "〇": 0,
    "一": 1,
    "二": 2,
    "两": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
}
CHINESE_UNITS = {"十": 10, "百": 100}


def chinese_number_to_int(text: str) -> int:
    """Convert the simple Chinese integers used before '台' to an int."""
    if text.isdigit():
        return int(text)

    total = 0
    current = 0
    for char in text:
        if char in CHINESE_DIGITS:
            current = CHINESE_DIGITS[char]
        elif char in CHINESE_UNITS:
            total += (current or 1) * CHINESE_UNITS[char]
            current = 0
        else:
            raise ValueError(f"不支持的中文数字：{text}")
    return total + current


def parse_material_detail(value: object) -> list[tuple[str, int | None]]:
    """Return one (model, count) item per model found in a detail cell."""
    if value is None:
        return []

    detail = str(value).strip()
    if not detail:
        return []

    pairs = [
        (match.group("model").strip(), chinese_number_to_int(match.group("count")))
        for match in PAIR_PATTERN.finditer(detail)
    ]
    if pairs:
        return pairs

    models = [match.group(0).strip() for match in MODEL_ONLY_PATTERN.finditer(detail)]
    if models:
        return [(model, None) for model in models]

    return [(detail, None)]


def find_source_sheet(workbook, requested_sheet: str | None):
    sheets = [workbook[requested_sheet]] if requested_sheet else workbook.worksheets
    for sheet in sheets:
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30)):
            for cell in row:
                if str(cell.value).strip() == HEADER_NAME:
                    return sheet, cell.row, cell.column
    location = f"工作表“{requested_sheet}”" if requested_sheet else "工作簿前 30 行"
    raise ValueError(f"未在{location}找到“{HEADER_NAME}”列")


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def split_workbook(input_path: Path, output_path: Path, sheet_name: str | None) -> tuple[int, int]:
    keep_vba = input_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(input_path, keep_vba=keep_vba)
    source, header_row, detail_column = find_source_sheet(workbook, sheet_name)

    if RESULT_SHEET_NAME in workbook.sheetnames:
        workbook.remove(workbook[RESULT_SHEET_NAME])
    result = workbook.create_sheet(RESULT_SHEET_NAME, 0)

    source_headers = [source.cell(header_row, column).value for column in range(1, source.max_column + 1)]
    result_headers = source_headers + ["台数"]
    for column, value in enumerate(result_headers, start=1):
        target = result.cell(1, column, value)
        if column <= source.max_column:
            copy_cell_style(source.cell(header_row, column), target)
        else:
            copy_cell_style(source.cell(header_row, detail_column), target)

    output_row = 2
    source_rows = 0
    result_rows = 0
    for row_number in range(header_row + 1, source.max_row + 1):
        values = [source.cell(row_number, column).value for column in range(1, source.max_column + 1)]
        if not any(value not in (None, "") for value in values):
            continue

        source_rows += 1
        items = parse_material_detail(values[detail_column - 1]) or [("", None)]
        for model, count in items:
            for column, value in enumerate(values, start=1):
                output_value = model if column == detail_column else value
                target = result.cell(output_row, column, output_value)
                copy_cell_style(source.cell(row_number, column), target)
            count_cell = result.cell(output_row, source.max_column + 1, count)
            copy_cell_style(source.cell(row_number, detail_column), count_cell)
            result.row_dimensions[output_row].height = source.row_dimensions[row_number].height
            output_row += 1
            result_rows += 1

    result.freeze_panes = "A2"
    result.auto_filter.ref = f"A1:{get_column_letter(len(result_headers))}{max(1, result_rows + 1)}"
    for column in range(1, source.max_column + 1):
        letter = get_column_letter(column)
        result.column_dimensions[letter].width = source.column_dimensions[letter].width
    result.column_dimensions[get_column_letter(source.max_column + 1)].width = 10

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return source_rows, result_rows


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="拆分 Excel 的“物料明细”，只保留机型并新增台数列。"
    )
    parser.add_argument("input", type=Path, help="输入的 .xlsx/.xlsm 文件")
    parser.add_argument("-o", "--output", type=Path, help="输出文件路径")
    parser.add_argument("--sheet", help="指定源工作表；默认自动查找“物料明细”列")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    input_path = args.input.expanduser().resolve()
    if not input_path.exists():
        raise SystemExit(f"输入文件不存在：{input_path}")
    if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SystemExit("仅支持 .xlsx 和 .xlsm 文件")

    output_path = (
        args.output.expanduser().resolve()
        if args.output
        else input_path.with_name(f"{input_path.stem}_台数已拆分{input_path.suffix}")
    )
    if output_path == input_path:
        raise SystemExit("输出文件不能与输入文件相同，以免覆盖原始数据")

    try:
        source_rows, result_rows = split_workbook(input_path, output_path, args.sheet)
    except (KeyError, PermissionError, ValueError) as exc:
        raise SystemExit(f"处理失败：{exc}") from exc

    print(f"处理完成：{source_rows} 条原始记录，生成 {result_rows} 条拆分记录")
    print(f"输出文件：{output_path}")


if __name__ == "__main__":
    main()
